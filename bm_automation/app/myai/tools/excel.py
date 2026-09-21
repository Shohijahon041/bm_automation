"""Excel Tool — Excel fayllar bilan ishlash."""

from __future__ import annotations

import os
import tempfile
from typing import Any

from . import BaseTool
from ...utils.logger import get_logger

log = get_logger("myai.tools.excel")


class ExcelTool(BaseTool):
    """Excel fayl yaratish va to'ldirish."""

    name = "excel"
    description = "Excel fayl yaratish, to'ldirish, formatlash (xlsx/csv)"

    async def execute(self, action: str = "", **kwargs) -> Any:
        actions = {
            "create_report": self._create_report,
            "create_from_data": self._create_from_data,
            "to_csv": self._to_csv,
        }
        handler = actions.get(action)
        if handler is None:
            raise ValueError(f"Noma'lum excel action: {action}")
        return await handler(**kwargs)

    async def _create_report(
        self, title: str = "", data: list[dict] = None,
        sheet_name: str = "", headers: list[str] = None,
        extra_sheets: list[dict] = None, save_path: str = "",
        **kwargs,
    ) -> dict:
        """Oddiy hisobot Excel fayli yaratish.

        title — fayl nomi va default sheet nomi.
        data — jadval qatorlari (dict ro'yxati).
        sheet_name — birinchi varaq nomi (default: title).
        headers — ustunlar tartibi (berilmasa dict kalitlari ishlatiladi).
        extra_sheets — qo'shimcha varaqlar: [{name, data, headers}].
        save_path — aniq fayl yo'li (berilmasa tempdir + title).
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
        except ImportError:
            raise RuntimeError("openpyxl o'rnatilmagan: pip install openpyxl")

        wb = Workbook()
        ws = wb.active
        ws.title = (sheet_name or title or "Hisobot")[:31] or "Hisobot"

        # Sarlavha
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4",
                                  fill_type="solid")
        header_font_white = Font(bold=True, size=12, color="FFFFFF")

        def _fill_sheet(sheet, rows: list[dict], cols: list[str]) -> int:
            if not rows:
                return 0
            if not cols:
                cols = list(rows[0].keys())
            for col, header in enumerate(cols, 1):
                cell = sheet.cell(row=1, column=col, value=header)
                cell.font = header_font_white
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
            for row_idx, row_data in enumerate(rows, 2):
                for col_idx, key in enumerate(cols, 1):
                    val = row_data.get(key)
                    if isinstance(val, (int, float)) and not isinstance(val, bool):
                        sheet.cell(row=row_idx, column=col_idx, value=val)
                    else:
                        sheet.cell(row=row_idx, column=col_idx, value=val)
            # Avtomatik kenglik
            for col in range(1, len(cols) + 1):
                max_len = max(
                    len(str(sheet.cell(row=r, column=col).value or ""))
                    for r in range(1, len(rows) + 2)
                )
                sheet.column_dimensions[chr(64 + col)].width = min(max_len + 2, 50)
            return len(rows)

        _fill_sheet(ws, data or [], headers or [])

        for extra in extra_sheets or []:
            name = (extra.get("name") or "")[:31] or "Qo'shimcha"
            if name in wb.sheetnames:
                continue
            sheet = wb.create_sheet(title=name)
            _fill_sheet(sheet, extra.get("data") or [],
                        extra.get("headers") or [])

        if not save_path:
            title_slug = "".join(
                c if c.isalnum() or c in " _-" else "_" for c in title
            ).strip().replace(" ", "_")
            save_path = os.path.join(
                tempfile.gettempdir(),
                f"myai_report_{title_slug or 'Hisobot'}.xlsx",
            )
        wb.save(save_path)
        log.info("Excel hisobot yaratildi: %s", save_path)
        return {"path": save_path, "rows": len(data or []),
                "sheets": [ws.title] + [e.get("name") for e in (extra_sheets or [])]}

    async def _create_from_data(self, headers: list[str] = None,
                                rows: list[list] = None, **kwargs) -> dict:
        """Formatlangan Excel fayl yaratish."""
        data = []
        if headers and rows:
            for row in rows:
                data.append(dict(zip(headers, row)))
        return await self._create_report(title="Hisobot", data=data)

    async def _to_csv(self, title: str = "Hisobot", data: list[dict] = None,
                      headers: list[str] = None, save_path: str = "",
                      **kwargs) -> dict:
        """Ma'lumotni CSV faylga yozadi."""
        import csv as _csv

        rows = data or []
        cols = headers or (list(rows[0].keys()) if rows else [])
        if not save_path:
            title_slug = "".join(
                c if c.isalnum() or c in " _-" else "_" for c in title
            ).strip().replace(" ", "_")
            save_path = os.path.join(
                tempfile.gettempdir(),
                f"myai_report_{title_slug or 'Hisobot'}.csv",
            )
        with open(save_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = _csv.writer(f)
            writer.writerow(cols)
            for row in rows:
                writer.writerow([row.get(c, "") for c in cols])
        log.info("CSV hisobot yaratildi: %s", save_path)
        return {"path": save_path, "rows": len(rows)}
