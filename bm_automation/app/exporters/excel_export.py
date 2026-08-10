"""Excel fayl yaratish: duty jadvali va oylik/yillik yig'ma.

openpyxl'ga asoslangan. Business logicdan ajratilgan — faqat berilgan
ma'lumotlarni chiroyli xlsx'ga aylantiradi.
"""

from __future__ import annotations

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

__all__ = ["duty_to_excel", "monthly_excel"]

DUTY_HEADERS = [
    "Sana", "Grafik", "Smena", "Haydovchi", "Jadval raqami",
    "Avtobus (ra'qam)", "Garaj raqami", "Avtobus modeli",
    "Boshlanish", "Tugash", "Qatnovlar soni", "Almashtirilgan",
]

DUTY_HEADER_KEYS = {
    "Sana": "date", "Grafik": "graphName", "Smena": "shiftName",
    "Haydovchi": "driverName", "Jadval raqami": "timeTableNumber",
    "Avtobus (ra'qam)": "plateNum", "Garaj raqami": "garageNumber",
    "Avtobus modeli": "vehicleModel", "Boshlanish": "startTime",
    "Tugash": "endTime", "Qatnovlar soni": "tripCount",
    "Almashtirilgan": "isReplaced",
}

DUTY_WIDTHS = [12, 8, 8, 30, 12, 14, 12, 22, 11, 11, 13, 14]

_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_THIN = Side(style="thin", color="B0B0B0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def duty_to_excel(rows: list[dict], out_file: str) -> str:
    """Duty qatorlarini Excel jadvaliga yozadi."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Duty"
    ws.append(DUTY_HEADERS)
    for r in rows:
        ws.append([r.get(DUTY_HEADER_KEYS[h]) for h in DUTY_HEADERS])
    for i, w in enumerate(DUTY_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(out_file)
    return out_file


def _cell(c, text, fill=None, bold=False):
    v = c.value = text
    c.alignment = Alignment(horizontal="center", vertical="center")
    if fill:
        c.fill = fill
    if bold:
        c.font = Font(bold=True)
    return c


def monthly_excel(rows: list[dict], graphs: list[str], driver_rows: list[dict],
                  out_file: str) -> str:
    """Oylik/yillik duty yig'masini Excel qilib yozadi.

    Sheet "Kunlik jadval": har kun x har grafik (haydovchi, avtobus).
    Sheet "Haydovchi xulosasi": haydovchi bo'yicha statistik.
    """
    wb = openpyxl.Workbook()

    # Sheet 1: Kunlik jadval
    ws = wb.active
    ws.title = "Kunlik jadval"
    headers = ["Sana", "Kun"]
    for g in graphs:
        headers += [f"{g}: Haydovchi", f"{g}: Avtobus"]
    ws.append(headers)
    for i, h in enumerate(headers, start=1):
        _cell(ws.cell(row=1, column=i), h, _HEADER_FILL, True)
        ws.cell(row=1, column=i).border = _BORDER

    for r in rows:
        rec = [r["date"], r["weekday"]]
        for g in graphs:
            c = r["graphs"].get(g, {"driver": "", "bus": ""})
            rec += [c["driver"], c["bus"]]
        ws.append(rec)
        for col in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=col).border = _BORDER

    widths = [12, 14]
    for _g in graphs:
        widths += [36, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Sheet 2: Haydovchi xulosasi
    ws2 = wb.create_sheet("Haydovchi xulosasi")
    h2 = ["Haydovchi", "Ishlagan kun", "Smena soatlari", "Avtobuslar", "Grafiklar"]
    ws2.append(h2)
    for i, h in enumerate(h2, start=1):
        _cell(ws2.cell(row=1, column=i), h, _HEADER_FILL, True)
        ws2.cell(row=1, column=i).border = _BORDER
    for r in driver_rows:
        ws2.append([r["driver"], r["days"], r["hours"], r["buses"], r["graphics"]])
        for col in range(1, len(h2) + 1):
            ws2.cell(row=ws2.max_row, column=col).border = _BORDER
    for i, w in enumerate([40, 12, 14, 24, 18], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    wb.save(out_file)
    return out_file
