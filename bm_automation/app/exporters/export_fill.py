"""Yo'nalish Jadvali Export xlsx'sini to'ldirish (faqat Excel ishi).

Saytdan olingan shift grafik export shabloni berilgan ma'lumotlar bilan
to'ldiriladi:
    - sarlavhada sana yoziladi
    - "Иш куни" katagiga haydovchi ism-familyasi yoziladi
    - grafik raqami (P1, P2...) avtobus raqami bilan yoziladi
    - tag qismiga mexanik BD va dispetcher nomlari qo'shiladi

Ma'lumot yig'ish (duty/shift API) `app/services/export_service` da.
"""

from __future__ import annotations

import os
import tempfile
import threading
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill

from .excel_cells import set_cell

__all__ = ["fill_export_workbook"]

_LUNCH_EMOJI = "\U0001F37D\ufe0f"


def fill_export_workbook(template_bytes: bytes, date_label: str,
                         by_graph: dict, out_path: str | Path,
                         graph_name: str = "",
                         mech_name: str = "",
                         dispatcher_name: str = "") -> Path:
    """Export shablonini to'ldirib, `out_path`ga saqlaydi.

    by_graph: {"P1": {"driver": str, "plate": str, "graph": str}, ...}
    graph_name: grafik nomi (masalan "B-80")
    mech_name: mexanik BD nomi
    dispatcher_name: dispetcher nomi
    """
    # ★ Thread-safe: har bir chaqiruv noyob vaqtinchalik fayl yaratadi
    fd, tmp_path_str = tempfile.mkstemp(suffix=".xlsx",
                                        prefix=f"graph_export_{threading.get_ident()}_")
    os.close(fd)
    tmp = Path(tmp_path_str)
    try:
        tmp.write_bytes(template_bytes)

        wb = openpyxl.load_workbook(tmp)
        ws = wb.active

        bold = Font(bold=True, size=12)
        bold_big = Font(bold=True, size=14)
        center = Alignment(horizontal="center", vertical="center")
        left = Alignment(horizontal="left", vertical="center")
        right = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2",
                                  fill_type="solid")
        gray_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2",
                                fill_type="solid")

        # Blok satrlarini topish — turli shablon formatlari:
        #   Format A (API): C1="Шанба куни"/"Иш куни", C6="1","2"...
        #   Format B (local): C1="Иш куни", C5="P1","P2"...
        #   Format C (B-80 local): C5="P1","P2" (kun nomi yo'q)
        #   Format D (test/simple): C1="Иш куни" (raqamsiz, ketma-ket indeks)
        _DAY_KEYWORDS = ("Иш куни", "Шанба куни", "Якшанба куни", "ISH KUNI",
                         "SHANBA KUNI", "YAKSHANBA KUNI")
        block_rows = []
        block_nums = []
        seq = 0
        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row,
                                                  max_col=min(ws.max_column, 8)),
                                     start=1):
            c1 = str(row[0].value or "").strip() if len(row) > 0 else ""
            c5 = str(row[4].value or "").strip() if len(row) > 4 else ""
            c6 = str(row[5].value or "").strip() if len(row) > 5 else ""
            num = None
            # Format A: kun nomi + raqam C6 (API shablon)
            if any(kw in c1 for kw in _DAY_KEYWORDS):
                if c6.isdigit():
                    num = int(c6)
                elif c5.upper().startswith("P") and c5[1:].isdigit():
                    num = int(c5[1:])
                else:
                    seq += 1
                    num = seq
            # Format C: faqat "P<n>" (local B-80)
            elif c5.upper().startswith("P") and c5[1:].isdigit():
                num = int(c5[1:])
            if num is not None:
                block_rows.append(r_idx)
                block_nums.append(num)

        print(f"  [export] Shablon: {ws.max_row} qator, {ws.max_column} ustun, "
              f"blocklar: {len(block_rows)}, block_nums: {block_nums}, "
              f"by_graph kalitlari: {list(by_graph.keys())}")

        for idx, r in enumerate(block_rows):
            pnum = block_nums[idx]
            info = by_graph.get(f"P{pnum}") or by_graph.get(str(pnum)) or {}
            driver = info.get("driver") or ""
            plate = info.get("plate") or ""
            graph_num = info.get("graph") or f"P{pnum}"

            # Sana (1-qator yuqorida)
            set_cell(ws, max(r - 1, 1), 1, date_label, bold, center)

            # Haydovchi ismi (2-ustun)
            set_cell(ws, r, 2, driver, bold, left)

            # Grafik nomeri (5-ustun) — "P1"
            set_cell(ws, r, 5, graph_num, bold, center)

            # Avtobus raqami (6-ustun) — "01A123AA"
            set_cell(ws, r, 6, plate, bold, center)

            # Bo'sh kataklarni to'ldirish — formatlash (3-5 ustunlarni tashlab ketish)
            for col in range(1, ws.max_column + 1):
                if col in (3, 4, 5):
                    continue
                cell = ws.cell(r, col)
                if cell.value is None or str(cell.value).strip() == "":
                    cell.border = thin_border

        # Tag qismiga mexanik BD + dispetcher qo'shish
        if block_rows:
            last_block_row = max(block_rows)
            footer_row = last_block_row + 3

            # Ajratuvchi chiziq
            for col in range(1, min(ws.max_column + 1, 9)):
                cell = ws.cell(footer_row, col)
                cell.border = Border(top=Side(style="double"))

            footer_row += 1
            if mech_name:
                set_cell(ws, footer_row, 1,
                         f"Mexanik BD: {mech_name}", bold, left)
                ws.merge_cells(
                    start_row=footer_row, start_column=1,
                    end_row=footer_row, end_column=3)

            if dispatcher_name:
                set_cell(ws, footer_row, 5,
                         f"Dispetcher: {dispatcher_name}", bold, right)
                ws.merge_cells(
                    start_row=footer_row, start_column=5,
                    end_row=footer_row, end_column=7)

            # Imzo qatori
            footer_row += 2
            if mech_name:
                set_cell(ws, footer_row, 1,
                         "___________________", None, center)
                set_cell(ws, footer_row + 1, 1,
                         "Mexanik BD", None, center)
                ws.merge_cells(
                    start_row=footer_row, start_column=1,
                    end_row=footer_row, end_column=3)
                ws.merge_cells(
                    start_row=footer_row + 1, start_column=1,
                    end_row=footer_row + 1, end_column=3)

            if dispatcher_name:
                set_cell(ws, footer_row, 5,
                         "___________________", None, center)
                set_cell(ws, footer_row + 1, 5,
                         "Dispetcher", None, center)
                ws.merge_cells(
                    start_row=footer_row, start_column=5,
                    end_row=footer_row, end_column=7)
                ws.merge_cells(
                    start_row=footer_row + 1, start_column=5,
                    end_row=footer_row + 1, end_column=7)

        # Ustun kengliklarini avtomatik moslashtirish
        from openpyxl.cell.cell import MergedCell
        from openpyxl.utils import get_column_letter
        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for row_idx in range(1, ws.max_row + 1):
                cell = ws.cell(row_idx, col_idx)
                if isinstance(cell, MergedCell):
                    continue
                try:
                    val = str(cell.value or "")
                    max_len = max(max_len, len(val))
                except Exception:
                    pass
            adjusted = min(max(max_len + 2, 8), 40)
            ws.column_dimensions[col_letter].width = adjusted

        # ● belgilarini tushlik emoji bilan almashtirish
        replaced = 0
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "\u25cf" in c.value:
                    c.value = c.value.replace("\u25cf", _LUNCH_EMOJI)
                    replaced += 1
        if replaced:
            print(f"  [export] {replaced} ta ● → 🍽️ almashtirildi")

        # Bosib chiqarish uchun landscape (albom) formati
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = 0.4
        ws.page_margins.right = 0.4
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5

        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out_path)
    finally:
        # ★ Har doim vaqtinchalik faylni tozalash
        try:
            tmp.unlink(missing_ok=True)
        except OSError:
            pass

    return out_path
