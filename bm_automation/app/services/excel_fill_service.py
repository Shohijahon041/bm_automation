"""Excel orqali saytning kunlik jadvalini (duty) to'ldirish xizmati.

Foydalanuvchi sayt export formatidagi Excel'ni yuboradi:
    Sana, Grafik, Smena, Haydovchi, Jadval raqami, Avtobus, Garaj raqami,
    Avtobus modeli, Boshlanish, Tugash, Qatnovlar soni, Almashtirilgan.

Bosqichlar:
  1. Excel o'qiladi (sana fayl ichidagi katakdan aniqlanadi).
  2. Haydovchi nomi -> driverId, avtobus raqami -> vehicleId API'dan topiladi.
  3. Saytga kunlik jadval yoziladi (POST /duty yoki PUT /duty/graph/replace).
  4. Natijada kunlik jadval rasmi + Excel qaytariladi.

ID yechish `app.repositories` (driver/vehicle), Excel chiqishi
`app.exporters` orqali.
"""

from __future__ import annotations

import re
from datetime import date, datetime, time
from pathlib import Path

import openpyxl

from ..api.client import BMClient
from ..config.settings import BRUTTO_MGMT
from ..repositories.driver_repo import DriverRepository
from ..repositories.duty_repo import DUTY, DutyRepository
from ..repositories.shift_repo import SHIFTS, ShiftRepository
from ..repositories.vehicle_repo import VehicleRepository
from ..utils.io import ensure_dir
from ..utils.text import normalize_uz, upper_norm

__all__ = [
    "DUTY",
    "DRIVERS",
    "SHIFTS",
    "VEHICLES",
    "build_payload",
    "fill_duty",
    "make_plan_report",
    "parse_plan_excel",
    "resolve_driver_id",
    "resolve_ids",
    "resolve_vehicle_id",
    "run",
]

DRIVERS = f"{BRUTTO_MGMT}/drivers"
VEHICLES = f"{BRUTTO_MGMT}/vehicles"

# Excel ustunlari -> ichki maydon
_HEADER_FIELDS = {
    "sana": "date",
    "grafik": "graph",
    "smena": "shift",
    "haydovchi": "driver",
    "jadval raqami": "time_table",
    "jadval": "time_table",
    "avtobus": "plate",
    "garaj raqami": "garage",
    "garaj": "garage",
    "avtobus modeli": "model",
    "model": "model",
    "boshlanish": "start",
    "tugash": "end",
    "qatnovlar soni": "trips",
    "almashtirilgan": "replaced",
}


def _norm(s) -> str:
    return normalize_uz(s)


def _norm_up(s) -> str:
    return upper_norm(s)


def _field_of(header: str) -> str | None:
    h = _norm(header).lower()
    if not h:
        return None
    if h in _HEADER_FIELDS:
        return _HEADER_FIELDS[h]
    for key, field in _HEADER_FIELDS.items():
        if h.startswith(key):
            return field
    return None


def _to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, time):
        return None
    s = _norm(v)
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    m = re.search(r"(\d{2})[./-](\d{2})[./-](\d{4})", s)
    if m:
        d, mo, y = map(int, m.groups())
        if 1 <= mo <= 12 and 1 <= d <= 31 and 2000 <= y <= 2100:
            return date(y, mo, d)
    return None


def _to_time(v) -> str:
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, datetime):
        return v.strftime("%H:%M")
    s = _norm(v)
    m = re.match(r"(\d{1,2}):(\d{2})", s)
    if m:
        return f"{int(m.group(1)):02d}:{m.group(2)}"
    m = re.match(r"(\d{1,2})\.(\d{2})", s)
    if m:
        return f"{int(m.group(1)):02d}:{m.group(2)}"
    return ""


def _find_header(ws):
    """Sarlavha qatorini topadi va (qator_indeks, {ustun_nomi: field}) qaytaradi."""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=12, max_col=20), start=1):
        fields = {}
        for cell in row:
            f = _field_of(cell.value)
            if f:
                fields[f] = cell.column
        if fields.get("driver") and (fields.get("graph") or fields.get("plate") or fields.get("start")):
            return i, fields
    return None, None


def _find_date(ws) -> date | None:
    for row in ws.iter_rows(min_row=1, max_row=15, max_col=20):
        for cell in row:
            d = _to_date(cell.value)
            if d:
                return d
    return None


def parse_plan_excel(path: str) -> dict:
    """Excel'ni o'qiydi. {date, rows:[{graph, shift, driver, plate, ...}]}."""
    wb = openpyxl.load_workbook(path, data_only=True)
    if not wb.sheetnames:
        raise ValueError("Excel bo'sh")
    for ws in wb.worksheets:
        header_row, fields = _find_header(ws)
        if not fields:
            continue
        sheet_date = _find_date(ws)
        rows = []
        for row in ws.iter_rows(min_row=header_row + 1, max_col=max(fields.values())):
            r = {}
            for f, col in fields.items():
                r[f] = row[col - 1].value
            graph = _norm(r.get("graph") or "")
            driver = _norm(r.get("driver") or "")
            if not graph and not driver:
                continue
            rows.append({
                "graph": graph,
                "shift": _norm(r.get("shift") or ""),
                "driver": driver,
                "time_table": r.get("time_table"),
                "plate": _norm(r.get("plate") or ""),
                "garage": _norm(r.get("garage") or ""),
                "model": _norm(r.get("model") or ""),
                "start": _to_time(r.get("start")),
                "end": _to_time(r.get("end")),
                "trips": r.get("trips"),
                "replaced": r.get("replaced"),
            })
        if not rows:
            continue
        day = sheet_date
        if not day:
            raise ValueError("Excel'da sana topilmadi (Sana ustuni yoki sana katagi bo'lishi kerak)")
        return {"date": day.isoformat(), "rows": rows}
    raise ValueError("Excel'da jadval qatorlari topilmadi (Grafik/Haydovchi ustunlari kerak)")


# ---- ID'lar ----

def resolve_driver_id(client: BMClient, route_variant_id: str, date_str: str, driver_name: str) -> str:
    return DriverRepository(client).resolve_driver_id(route_variant_id, date_str, driver_name)


def resolve_vehicle_id(client: BMClient, route_variant_id: str, date_str: str,
                       plate: str, garage: str = "") -> str:
    return VehicleRepository(client).resolve_vehicle_id(route_variant_id, date_str, plate, garage)


def resolve_ids(client: BMClient, route_variant_id: str, date_str: str, rows: list[dict]) -> dict:
    """Barcha noyob haydovchi/avtobuslar uchun ID'larni topadi (cache bilan)."""
    driver_repo = DriverRepository(client)
    vehicle_repo = VehicleRepository(client)
    drivers = {}
    for r in rows:
        key = upper_norm(r.get("driver") or "")
        if key and key not in drivers:
            drivers[key] = driver_repo.resolve_driver_id(route_variant_id, date_str, key)
    vehicles = {}
    for r in rows:
        key = (upper_norm(r.get("plate") or ""), upper_norm(r.get("garage") or ""))
        if key not in vehicles:
            vehicles[key] = vehicle_repo.resolve_vehicle_id(
                route_variant_id, date_str, key[0], key[1])
    return {"drivers": drivers, "vehicles": vehicles}


# ---- Template va payload ----

def _shift_template(client: BMClient, route_variant_id: str, date_str: str) -> dict:
    """shiftId + shiftGraphId ro'yxati (GET /duty yoki GET /shifts/by-date orqali)."""
    duty_repo = DutyRepository(client)
    shift_repo = ShiftRepository(client)
    duty = None
    try:
        duty = duty_repo.by_date(route_variant_id, date_str)
    except Exception:
        duty = None
    if isinstance(duty, dict) and duty.get("graphs"):
        return {"duty": duty, "shiftId": duty.get("shiftId") or "",
                "shiftName": duty.get("shiftName") or "",
                "graphs": duty["graphs"]}
    try:
        sh = shift_repo.by_date(route_variant_id, date_str)
        if isinstance(sh, dict) and sh.get("graphicList"):
            shift_name = sh.get("name") or ""
            return {
                "duty": None,
                "shiftId": sh.get("id") or "",
                "shiftName": shift_name,
                "graphs": [{"shiftGraphId": g.get("id"), "graphName": g.get("name"),
                            "shiftName": shift_name} for g in sh["graphicList"]],
            }
    except Exception:
        pass
    if isinstance(duty, dict) and duty.get("shiftId"):
        return {"duty": duty, "shiftId": duty["shiftId"],
                "shiftName": duty.get("shiftName") or "",
                "graphs": duty.get("graphs") or []}
    raise ValueError("Kunlik jadval template'i topilmadi (shift ma'lumoti yo'q)")


def build_payload(template: dict, rows: list[dict], ids: dict) -> list[dict]:
    """Excel qatorlaridan createDuty/replace graphs ro'yxatini quradi."""
    by_graph = {}
    for r in rows:
        key = upper_norm(r.get("graph") or "")
        by_graph.setdefault(key, []).append(r)

    graphs_out = []
    for g in template["graphs"]:
        gn = upper_norm(g.get("graphName") or "")
        entries = by_graph.get(gn, [])
        main = entries[0] if entries else {}
        second = entries[1] if len(entries) > 1 else {}

        plate = upper_norm(main.get("plate") or "")
        garage = upper_norm(main.get("garage") or "")
        did = ids["drivers"].get(upper_norm(main.get("driver") or "")) or ""
        sdid = ids["drivers"].get(upper_norm(second.get("driver") or "")) or ""
        vid = ids["vehicles"].get((plate, garage)) or ""

        trips = main.get("trips")
        try:
            trips = int(float(trips)) if trips not in (None, "") else None
        except (ValueError, TypeError):
            trips = None

        graphs_out.append({
            "shiftGraphId": g.get("shiftGraphId") or "",
            "graphName": g.get("graphName") or main.get("graph") or "",
            "shiftName": g.get("shiftName") or main.get("shift") or "",
            "driverId": did,
            "vehicleId": vid,
            "startTime": main.get("start") or "",
            "endTime": main.get("end") or "",
            "timeTableNumber": str(main.get("time_table") or ""),
            "tripCount": trips,
            "hasSecond": bool(sdid),
            "secondDriverId": sdid,
            "secondStartTime": second.get("start") or "",
            "secondEndTime": second.get("end") or "",
            "secondStartDir": "UP",
        })
    return graphs_out


# ---- Yozish ----

def fill_duty(client: BMClient, route_variant_id: str, date_str: str,
              rows: list[dict], dry_run: bool = False) -> dict:
    """Saytga kunlik jadvalni yozadi (POST /duty yoki PUT /duty/graph/replace)."""
    duty_repo = DutyRepository(client)
    template = _shift_template(client, route_variant_id, date_str)

    existing_filled = any(
        g.get("driverId") for g in (template.get("duty") or {}).get("graphs", []))
    if existing_filled:
        raise ValueError(
            "Bu kunga jadval allaqachon to'ldirilgan. "
            "Qayta yozish uchun avval saytdagi jadvalni o'chiring yoki boshqa kun tanlang.")

    ids = resolve_ids(client, route_variant_id, date_str, rows)
    graphs = build_payload(template, rows, ids)

    missing = []
    seen_drivers, seen_vehicles = set(), set()
    for g in graphs:
        gname = g.get("graphName") or "?"
        did = g.get("driverId") or ""
        vid = g.get("vehicleId") or ""
        if not did:
            missing.append(f"{gname}: haydovchi topilmadi")
        elif did in seen_drivers:
            missing.append(f"{gname}: haydovchi dublikat")
        seen_drivers.add(did)
        if not vid:
            missing.append(f"{gname}: avtobus topilmadi")
        elif vid in seen_vehicles:
            missing.append(f"{gname}: avtobus dublikat")
        seen_vehicles.add(vid)
    if missing:
        raise ValueError(
            "Jadval yozilmadi: sayt barcha grafiklar uchun noyob haydovchi va avtobus "
            "talab qiladi. Excel'da quyidagilar yetishmayapti:\n  " + "\n  ".join(missing))

    payload = {
        "routeVariantId": route_variant_id,
        "date": date_str,
        "shiftId": template["shiftId"],
        "graphs": graphs,
    }

    if dry_run:
        return {"dry_run": True, "payload": payload, "ids": ids}

    duty_id = (template.get("duty") or {}).get("id")
    try:
        if duty_id:
            client.put(f"{DUTY}/graph/replace", json={**payload, "dutyId": duty_id})
        else:
            client.post(DUTY, json=payload)
    except Exception as exc:
        # yangi duty bo'lsa, mavjud bo'lib chiqsa POST o'rniga replace
        if "duty_id" in str(exc).lower() or "ziddiyat" in str(exc).lower() or "mavjud" in str(exc).lower():
            saved_duty = duty_repo.by_date(route_variant_id, date_str)
            if isinstance(saved_duty, dict) and saved_duty.get("id"):
                client.put(f"{DUTY}/graph/replace",
                           json={**payload, "dutyId": saved_duty["id"]})
            else:
                raise
        else:
            raise

    saved = duty_repo.by_date(route_variant_id, date_str)
    return {"payload": payload, "saved": saved, "ids": ids}


def make_plan_report(client: BMClient, route_id: str, date_str: str,
                     profile: dict | None, out_dir: str = "reports") -> dict:
    """To'ldirilgandan keyin rasm + Excel yaratadi."""
    from ..exporters.excel_export import duty_to_excel
    from ..services.driver_sheet_service import run as driver_sheet_run
    from ..services.duty_service import graphs_to_rows

    base = Path(out_dir) / "plan"
    ensure_dir(base)

    img_res = driver_sheet_run(
        client, route_id, date_str=date_str, out_dir=str(base), send=False, profile=profile)

    saved = DutyRepository(client).by_date(route_id, date_str)
    rows_x = graphs_to_rows(saved)
    xls = duty_to_excel(rows_x, str(base / f"duty_{date_str.replace('-', '')}.xlsx"))

    return {"image": img_res["image"], "excel": xls}


def run(
    client: BMClient,
    route_id: str,
    date_str: str,
    rows: list[dict],
    profile: dict | None = None,
    dry_run: bool = False,
    out_dir: str = "reports",
) -> dict:
    """To'liq jarayon: saytga yozish + rasm + Excel. {image, excel, saved, ids}."""
    res = fill_duty(client, route_id, date_str, rows, dry_run=dry_run)
    if dry_run:
        return res
    report = make_plan_report(client, route_id, date_str, profile, out_dir=out_dir)
    return {**res, **report}
