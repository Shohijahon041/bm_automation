"""Excel o'qish (rejadagi jadval) va export to'ldirish testlari."""

from datetime import date

import openpyxl
import pytest

from bm_automation.app.exporters.export_fill import fill_export_workbook
from bm_automation.app.services.excel_fill_service import (
    _to_date,
    _to_time,
    parse_plan_excel,
)


def _make_plan_xlsx(tmp_path, date_cell="01.08.2026"):
    path = tmp_path / "plan.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = f"10-сонли автобус йўналиши  {date_cell}"
    ws["A2"] = "Sana"
    ws["B2"] = "Grafik"
    ws["C2"] = "Smena"
    ws["D2"] = "Haydovchi"
    ws["E2"] = "Jadval raqami"
    ws["F2"] = "Avtobus"
    ws["G2"] = "Boshlanish"
    ws["H2"] = "Tugash"
    ws["A3"] = date_cell
    ws["B3"] = "P1"
    ws["C3"] = "DU"
    ws["D3"] = "Aliyev Aliy"
    ws["E3"] = "101"
    ws["F3"] = "01A123AA"
    ws["G3"] = "06:00"
    ws["H3"] = "14:00"
    ws["B4"] = "P2"
    ws["D4"] = "Karimov Karim"
    wb.save(path)
    return path


def test_parse_plan_excel(tmp_path):
    parsed = parse_plan_excel(str(_make_plan_xlsx(tmp_path)))
    assert parsed["date"] == "2026-08-01"
    assert len(parsed["rows"]) == 2
    first = parsed["rows"][0]
    assert first["graph"] == "P1"
    assert first["driver"] == "Aliyev Aliy"  # normalize_uz (katta-kichik saqlanadi)
    assert first["start"] == "06:00"
    assert first["end"] == "14:00"
    assert first["plate"] == "01A123AA"


def test_parse_plan_excel_missing_headers(tmp_path):
    path = tmp_path / "bad.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Hech qanday ustun yo'q"
    wb.save(path)
    with pytest.raises(ValueError):
        parse_plan_excel(str(path))


def test_to_date_parses_formats():
    assert _to_date("01.08.2026") == date(2026, 8, 1)
    assert _to_date("2026-08-01") == date(2026, 8, 1)
    assert _to_date("") is None


def test_to_time_parses():
    assert _to_time("06:00") == "06:00"
    assert _to_time("6.00") == "06:00"
    assert _to_time("") == ""


def _make_export_template(tmp_path):
    path = tmp_path / "template.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "10-сонли автобус йўналиши"
    ws["A3"] = "Иш куни"
    ws["B3"] = "Haydovchi"
    ws["F3"] = "Avtobus"
    ws["A13"] = "Иш куни"
    wb.save(path)
    return path


def test_fill_export_workbook(tmp_path):
    template = _make_export_template(tmp_path)
    by_graph = {
        "P1": {"driver": "Aliyev Aliy", "plate": "01A123AA", "graph": "P1"},
        "P2": {"driver": "Karimov Karim", "plate": "01B456BB", "graph": "P2"},
    }
    out = tmp_path / "filled.xlsx"
    result = fill_export_workbook(template.read_bytes(), "01.08.2026", by_graph, out)

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert ws["A2"].value == "01.08.2026"
    assert ws["B3"].value == "Aliyev Aliy"
    assert ws["E3"].value == "P1"
    assert ws["F3"].value == "01A123AA"
    assert ws["B13"].value == "Karimov Karim"
    assert ws["E13"].value == "P2"
    assert ws["F13"].value == "01B456BB"
    assert result == out


def _make_merged_export_template(tmp_path):
    """Sanasi birlashtirilgan (MergedCell) katak ichiga tushadigan shablon."""
    path = tmp_path / "merged_template.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "10-сонли автобус йўналиши"
    ws.merge_cells("A2:A3")          # A2 anchor, A3 -> MergedCell
    ws["A4"] = "Иш куни"
    ws["B4"] = "Haydovchi"
    ws["F4"] = "Avtobus"
    ws.merge_cells("A5:A6")          # A5 anchor, A6 -> MergedCell
    ws["A7"] = "Иш куни"
    wb.save(path)
    return path


def test_fill_export_workbook_merged_cells(tmp_path):
    """MergedCell ichiga yozish xato bermasligi, anchor katakka yozilishi."""
    template = _make_merged_export_template(tmp_path)
    by_graph = {"P1": {"driver": "Aliyev Aliy", "plate": "01A123AA", "graph": "P1"}}
    out = tmp_path / "filled_merged.xlsx"
    fill_export_workbook(template.read_bytes(), "02.08.2026", by_graph, out)

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert ws["A2"].value == "02.08.2026"   # A3 (MergedCell) -> anchor A2
    assert ws["B4"].value == "Aliyev Aliy"
    assert ws["E4"].value == "P1"
    assert ws["F4"].value == "01A123AA"


def test_set_cell_merged_anchor(tmp_path):
    from bm_automation.app.exporters.excel_cells import set_cell

    path = tmp_path / "m.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("B2:D4")
    set_cell(ws, 3, 3, "X")
    assert ws["B2"].value == "X"
    assert ws["C3"].value is None


def test_fill_grafik_merged_cells(tmp_path):
    from datetime import date as _date

    from daily_grafik import fill_grafik

    tpl = tmp_path / "graf.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["E2"] = "P1"
    ws["E4"] = "P2"
    ws.merge_cells("A1:C1")           # sarlavha birlashtirilgan
    ws["A1"] = "old"
    ws.merge_cells("F2:F3")           # avtobus raqami hujayrasi merged
    ws["F2"] = "old-plate"
    ws.merge_cells("A4:B4")
    wb.save(tpl)

    out = tmp_path / "out.xlsx"
    by_graph = {"P1": {"driver": "Aliyev Aliy", "plate": "01A123AA"}}
    title, nblocks, filled, _ = fill_grafik(
        str(tpl), str(out), "B-80", "ISH KUNI", _date(2026, 8, 17), by_graph)

    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    assert nblocks == 2
    assert filled == 1
    assert ws2["A1"].value == title        # merged sarlavhaga yozildi
    assert ws2["A2"].value == "Aliyev Aliy"
    assert ws2["F2"].value == "01A123AA"   # F3 MergedCell -> anchor F2


def test_bot_targets_bot_only(monkeypatch):
    """`_bot_targets`: grafik FAQAT bot ichida — guruhlarga hech qachon.

    Qoida: shaxsiy chat_id -> o'sha chat; aks holda (chat_id yo'q yoki
    guruh) -> ADMIN/DISPATCHER chatlari (`grafik_sms._targets`).
    """
    from bm_automation.app.notifications.ops import grafik_sms as gs
    from daily_grafik import _bot_targets

    monkeypatch.setattr(gs, "_targets", lambda: ["-100main", "-100adm"])
    # Shaxsiy chat (tugma bosilgan joy) -> faqat shu chat.
    assert _bot_targets("r1", "777") == ["777"]
    assert _bot_targets("r1", 12345) == ["12345"]
    # chat_id yo'q (avtomatik oqim) -> ADMIN/DISPATCHER chatlari.
    assert _bot_targets("r1", None) == ["-100main", "-100adm"]
    # Guruh chat_id -> guruhga EMAS, ADMIN/DISPATCHER'ga.
    assert _bot_targets("r1", "-100GROUP") == ["-100main", "-100adm"]
    # ADMIN/DISPATCHER ham sozlanmagan -> [] (hech qayerga yuborilmaydi).
    monkeypatch.setattr(gs, "_targets", lambda: [])
    assert _bot_targets("r1", None) == []
    assert _bot_targets("r1", "-100GROUP") == []


def test_grafik_sent_once_per_day(tmp_path, monkeypatch):
    """Excel+rasm va eslatmalar kuniga bir marta — holat fayli bilan."""
    import daily_grafik as dg
    monkeypatch.setattr(dg, "STATE_DIR", tmp_path / "state")
    rid, ds = "r-1", "2026-09-30"
    assert dg._grafik_was_sent(rid, ds) is False
    # "e'lon qilinmagan" eslatmasi bir marta yuboriladi.
    assert dg._grafik_notify_once(rid, ds, "notfound") is True
    assert dg._grafik_notify_once(rid, ds, "notfound") is False
    assert dg._grafik_notify_once(rid, ds, "error") is True
    # Excel+rasm belgilangach takroriy jo'natish bloklanadi.
    dg._mark_grafik_sent(rid, ds)
    assert dg._grafik_was_sent(rid, ds) is True
    assert dg._grafik_was_sent(rid, "2026-10-01") is False


def test_process_route_skips_when_already_sent(monkeypatch):
    """Allaqachon yuborilgan (yo'nalish, kun) — duty ham olinmaydi, qayta yuborilmaydi."""
    import datetime

    import daily_grafik as dg
    calls = {"duty": 0}

    def boom(*a, **k):
        calls["duty"] += 1
        raise AssertionError("duty olinmasligi kerak")

    monkeypatch.setattr(dg, "_grafik_was_sent", lambda rid, ds: True)
    monkeypatch.setattr(dg, "_duty_with_retry", boom)
    res = dg._process_route(
        ("", "r1", "img", "B-80", 0), None, datetime.date(2026, 9, 30),
        {}, None, "ISH KUNI", retry_until_published=True)
    assert calls["duty"] == 0
    assert res == [("already", "B-80", True)]


def test_process_route_duplicate_flow_waits(monkeypatch):
    """Ikkinchi so'rov kutish oqimini takror ochmaydi — ('waiting') qaytadi."""
    import datetime
    import threading

    import daily_grafik as dg
    from bm_automation.app.notifications import sms_notify

    entered = threading.Event()
    release = threading.Event()
    route = ("", "r1", "img", "B-80", 0)
    day = datetime.date(2026, 10, 2)

    def fake_retry(client, rid, d, label, on_wait=None):
        entered.set()
        release.wait(timeout=5)
        return {"graphs": []}

    def boom(*a, **k):
        raise RuntimeError("test")

    monkeypatch.setattr(dg, "_grafik_was_sent", lambda rid, ds: False)
    monkeypatch.setattr(dg, "_duty_with_retry", fake_retry)
    monkeypatch.setattr(dg, "_bot_targets", lambda rid, cid: [])
    monkeypatch.setattr(sms_notify, "send_sms_for_graphs",
                        lambda *a, **k: {"sent": 0, "skipped": 0, "failed": 0,
                                         "invalid": 0, "removed": 0,
                                         "errors": 0})
    monkeypatch.setattr("bm_automation.app.services.export_service.build_export",
                        boom)
    monkeypatch.setattr("bm_automation.app.services.driver_sheet_service.run",
                        boom)
    dg._INFLIGHT.clear()

    t = threading.Thread(target=dg._process_route,
                         args=(route, None, day, {}, None, "ISH KUNI"))
    t.start()
    try:
        assert entered.wait(timeout=5)
        # Birinchi oqim kutayotganda ikkinchi so'rov yangi oqim ochmaydi.
        res = dg._process_route(route, None, day, {}, None, "ISH KUNI",
                                retry_until_published=True)
        assert res == [("waiting", "B-80", True)]
    finally:
        release.set()
    t.join(timeout=10)
    assert not t.is_alive()
    # Oqim tugagach band bo'shatiladi — keyingi urinishlarga xalaqit bermaydi.
    assert "r1|2026-10-02" not in dg._INFLIGHT


def test_duty_with_retry_calls_on_wait(monkeypatch):
    """404 bo'lganda on_wait chaqiriladi, keyin qayta urinib qaytaradi."""
    import datetime

    import daily_grafik as dg

    calls = {"n": 0, "wait": 0}

    class FakeRepo:
        def __init__(self, client):
            pass

        def by_date(self, rid, ds):
            calls["n"] += 1
            if calls["n"] == 1:
                raise RuntimeError("hali yo'q")
            return {"graphs": []}

    monkeypatch.setattr(dg, "DutyRepository", FakeRepo)
    monkeypatch.setattr(dg, "_not_found",
                        lambda exc: isinstance(exc, RuntimeError))
    monkeypatch.setattr(dg.time, "sleep", lambda s: None)
    monkeypatch.setattr(dg, "_warn_if_late", lambda rid, d, label: None)

    def on_wait():
        calls["wait"] += 1

    res = dg._duty_with_retry(None, "r1", datetime.date(2026, 10, 2), "B-80",
                              on_wait=on_wait)
    assert res == {"graphs": []}
    assert calls["n"] == 2
    assert calls["wait"] == 1


def test_summary_groups_and_text():
    """Xulosa: sent/already/waiting/notfound/error to'g'ri ajratiladi."""
    import datetime

    import daily_grafik as dg

    results = [("already", "B-80", True), ("waiting", "D-4", True),
               ("excel", "B-7", True), ("image", "B-7", True),
               ("excel", "E-17", True), ("image", "E-17", False),
               ("notfound", "10-YO'NALISH", True),
               ("route", "B-11", False)]
    s = dg._summarize_results(results)
    assert s["sent"] == ["B-7"]
    assert s["already"] == ["B-80"]
    assert s["waiting"] == ["D-4"]
    assert set(s["error"]) == {"E-17", "B-11"}
    assert s["notfound"] == ["10-YO'NALISH"]

    text = dg._summary_text(s, datetime.date(2026, 10, 2), "Ertaga")
    assert "02.10.2026" in text
    assert "✅ Yuborildi: B-7" in text
    assert "📤 Allaqachon yuborilgan: B-80" in text
    assert "⏳" in text and "D-4" in text
    assert "❌ Xatolik: B-11, E-17" in text or "❌ Xatolik: E-17, B-11" in text


def test_summary_empty():
    """Natija bo'sh bo'lsa ham xulosa yaroqli bo'lishi kerak."""
    import datetime

    import daily_grafik as dg

    text = dg._summary_text(dg._summarize_results([]),
                            datetime.date(2026, 10, 2), "Ertaga")
    assert "Ma'lumot yo'q" in text
