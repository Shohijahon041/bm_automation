"""Dashboard ko'p korxonali login/sessiya va rol qamrovi testlari."""

import io
import json

import pytest

from bm_automation.app.core import profiles as profiles_mod
from bm_automation.app.dashboard import server as server_mod
from bm_automation.app.db.storage import storage_for
from bm_automation.app.dashboard.server import DashboardHandler
from tests.sqlite_backend import SQLiteDatabase


class _FakeConn:
    def __init__(self, body=b""):
        self._body = body
        self._wbuf = io.BytesIO()

    def makefile(self, mode, *args):
        return io.BytesIO(self._body) if mode.startswith("r") else self._wbuf

    def sendall(self, data):
        self._wbuf.write(data)


def _request(storage, path, method="GET", headers=None, body=None,
             monkeypatch=None):
    """DashboardHandler'ni to'g'ridan-to'g'ri chaqirib natijani qaytaradi.

    `BM_TEST_MODE` o'chirilgan bo'lishi kerak (haqiqiy auth tekshiriladi).
    Qaytaradi: (status, json, cookies).
    """
    conn = _FakeConn(body or b"")
    h = object.__new__(DashboardHandler)
    h.server = type("S", (), {"daemon_threads": True})()
    h.client_address = ("127.0.0.1", 0)
    h.command = method
    h.path = path
    h.requestline = f"{method} {path} HTTP/1.0"
    h.request_version = "HTTP/1.0"
    h.headers = headers or {}
    h.rfile = io.BytesIO(body or b"")
    h.wfile = conn._wbuf
    if method == "GET":
        h.do_GET()
    else:
        h.do_POST()
    raw = conn._wbuf.getvalue()
    head, _, payload = raw.partition(b"\r\n\r\n")
    cookies = {}
    for line in head.decode("latin1").split("\r\n"):
        if line.lower().startswith("set-cookie:"):
            for part in line.split(":", 1)[1].strip().split(";"):
                kv = part.strip().split("=", 1)
                if len(kv) == 2:
                    cookies[kv[0]] = kv[1]
    status = int(head.split()[1]) if head else 0
    data = json.loads(payload.decode("utf-8")) if payload else {}
    return status, data, cookies


@pytest.fixture()
def storage(tmp_path):
    db = SQLiteDatabase(str(tmp_path / "st.db"))
    return storage_for(db)


@pytest.fixture()
def env(monkeypatch, tmp_path):
    """Haqiqiy auth uchun muhit: BM_TEST_MODE o'chirilgan, sessiya tmp'da."""
    monkeypatch.delenv("BM_TEST_MODE", raising=False)
    monkeypatch.setattr(server_mod, "_DASHBOARD_TOKEN", "")
    monkeypatch.setattr(server_mod, "_sessions_file",
                        lambda: tmp_path / "sessions.json")
    with server_mod._SESSION_LOCK:
        server_mod._sessions = {}
    with server_mod._POST_LOCK:
        server_mod._last_post_at.clear()
    monkeypatch.setattr(server_mod, "_POST_MIN_INTERVAL_S", 0.0)
    db = SQLiteDatabase(str(tmp_path / "dash.db"))
    st = storage_for(db)
    monkeypatch.setattr("bm_automation.app.db.storage.get_storage", lambda: st)
    monkeypatch.setattr("bm_automation.app.dashboard.metrics.get_storage",
                        lambda: st)
    # Test kompaniyasi: XTEST -> route r1
    monkeypatch.setattr(profiles_mod, "PROFILES_FILE", tmp_path / "profiles.json")
    db2 = SQLiteDatabase(str(tmp_path / "profiles_seed.db"))
    _st2 = storage_for(db2)
    _st2.save_route(external_id="r1", name="10-yo'nalish", entity_type="ROUTE")
    _st2.save_route(external_id="r2", name="20-yo'nalish", entity_type="ROUTE")
    profiles_mod.save_profiles({
        "active": "XTEST",
        "profiles": [
            {"name": "XTEST", "routeVariantId": "r1",
             "routeName": "10-yo'nalish",
             "profileId": "", "start1": "", "start2": ""},
            {"name": "YTEST", "routeVariantId": "r2",
             "routeName": "20-yo'nalish",
             "profileId": "", "start1": "", "start2": ""},
        ],
    })
    return st


def _login(env, username, password):
    body = json.dumps({"username": username, "password": password}).encode()
    headers = {"Content-Length": str(len(body))}
    return _request(env, "/api/login", method="POST", headers=headers, body=body)


# ---------------------------------------------------------------- auth

def test_schema_creates_dashboard_users(storage):
    """Storagema dashboard_users jadvali mavjud va ishlaydi."""
    uid = storage.dashboard_user_add("ali", "hash", "salt", "MANAGER",
                                     "XTEST", 1)
    assert uid > 0
    row = storage.dashboard_user_get(username="ali")
    assert row["role"] == "MANAGER"
    assert row["company"] == "XTEST"
    assert row["salt"] == "salt"
    assert storage.dashboard_users_count() == 1


def test_dashboard_user_unique_username(storage):
    """Bir xil username ikki marta qo'shilmaydi (UNIQUE)."""
    storage.dashboard_user_add("ali", "h1", "s1")
    storage.dashboard_user_add("ali", "h2", "s2")
    rows = storage.dashboard_users_list()
    assert len(rows) == 1
    assert rows[0]["salt"] == "s1"


def test_dashboard_user_update_and_delete(storage):
    uid = storage.dashboard_user_add("ali", "hash", "salt")
    assert storage.dashboard_user_update(uid, role="DIRECTOR",
                                         company="XTEST", active=0)
    row = storage.dashboard_user_get(row_id=uid)
    assert row["role"] == "DIRECTOR"
    assert row["company"] == "XTEST"
    assert row["active"] == 0
    assert storage.delete_dashboard_user(uid)
    assert storage.dashboard_user_get(row_id=uid) is None


def test_password_hash_and_verify(env):
    """Parol salt bilan xeshlanadi va to'g'ri tekshiriladi."""
    digest, salt = server_mod._hash_password("sirli-parol")
    assert digest and salt and digest != "sirli-parol"
    assert server_mod._verify_password("sirli-parol", salt, digest)
    assert not server_mod._verify_password("boshqa", salt, digest)
    # Turli salt -> turli hash
    digest2, salt2 = server_mod._hash_password("sirli-parol")
    assert digest2 != digest


def test_login_requires_auth(env, monkeypatch):
    """Ruxsatsiz API so'rovi 401 qaytaradi."""
    status, data, _ = _request(env, "/api/summary", monkeypatch=monkeypatch)
    assert status == 401
    assert "Kirish" in data.get("error", "")


def test_login_wrong_password(env):
    st = env
    st.dashboard_user_add("ali", *server_mod._hash_password("to'g'ri"),
                          role="ADMIN")
    status, data, _ = _login(env, "ali", "noto'g'ri")
    assert status == 401
    assert data.get("ok") is False


def test_login_success_and_session(env):
    st = env
    st.dashboard_user_add("ali", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    status, data, cookies = _login(env, "ali", "to'g'ri")
    assert status == 200 and data["ok"] is True
    assert data["role"] == "DISPATCHER"
    assert data["company"] == "XTEST"
    token = cookies.get("bm_session")
    assert token
    # Sessiya orqali /api/me ishlaydi
    status, data, _ = _request(env, "/api/me",
                               headers={"Cookie": f"bm_session={token}"})
    assert status == 200 and data["ok"] is True
    assert data["username"] == "ali"
    assert data["role"] == "DISPATCHER"
    assert data["company"]["name"] == "XTEST"


def test_logout_destroys_session(env):
    st = env
    st.dashboard_user_add("ali", *server_mod._hash_password("to'g'ri"),
                          role="ADMIN")
    _, _, cookies = _login(env, "ali", "to'g'ri")
    token = cookies.get("bm_session")
    hdr = {"Cookie": f"bm_session={token}"}
    status, data, _ = _request(env, "/api/logout", method="POST",
                               headers=hdr, body=b"{}")
    assert status == 200 and data["ok"] is True
    status, _, _ = _request(env, "/api/me", headers=hdr)
    assert status == 401


def test_token_bearer_restores_admin(env, monkeypatch):
    """DASHBOARD_TOKEN bearer'da ADMIN qaytadi (DISPATCHER ham emas)."""
    from bm_automation.app.dashboard import server as _srv
    monkeypatch.setattr(_srv, "_DASHBOARD_TOKEN", "master-token")
    status, data, _ = _request(
        env, "/api/routes",
        headers={"Authorization": "Bearer master-token"})
    assert status == 200
    ids = [r.get("id") for r in data["routes"]]
    assert {"r1", "r2"} <= set(ids)


def test_session_wins_over_master_token(env, monkeypatch):
    """Sessiya bilan kirilgan foydalanuvchi master token ustidan o'zi bo'ladi.

    Loopback'da sahifaga in'ektsiya qilingan DASHBOARD_TOKEN har so'rovda
    Authorization header orqali keladi; bu token sessiyani bosib, DISPATCHER
    foydalanuvchini ADMIN qilib ko'rsatmasligi kerak.
    """
    from bm_automation.app.dashboard import server as _srv
    monkeypatch.setattr(_srv, "_DASHBOARD_TOKEN", "master-token")
    st = env
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}",
           "Authorization": "Bearer master-token"}
    # /api/me sessiya rolini qaytaradi, ADMIN emas.
    status, data, _ = _request(env, "/api/me", headers=hdr)
    assert status == 200 and data["ok"] is True
    assert data["role"] == "DISPATCHER"
    assert data["is_admin"] is False
    # Route'lar ham faqat o'z korxonasi.
    status, data, _ = _request(env, "/api/routes", headers=hdr)
    assert [r.get("id") for r in data["routes"]] == ["r1"]
    # ADMIN-only endpoint hali ham 403.
    status, _, _ = _request(env, "/api/users", headers=hdr)
    assert status == 403
    # Master token sessiyasiz bo'lsa ADMIN bo'lib qoladi.
    status, data, _ = _request(
        env, "/api/me",
        headers={"Authorization": "Bearer master-token"})
    assert status == 200
    assert data["is_admin"] is True


def test_inactive_user_cannot_login(env):
    st = env
    uid = st.dashboard_user_add("ali", *server_mod._hash_password("to'g'ri"),
                                role="DISPATCHER", company="XTEST")
    st.dashboard_user_update(uid, active=0)
    status, _, _ = _login(env, "ali", "to'g'ri")
    assert status == 401


# ---------------------------------------------------------------- scope

def test_routes_only_own_company(env, monkeypatch):
    """DISPATCHER faqat o'z korxonasi route'ini ko'radi."""
    st = env
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, data, _ = _request(env, "/api/routes", headers=hdr)
    assert status == 200
    ids = [r.get("id") for r in data["routes"]]
    assert ids == ["r1"]
    assert "r2" not in ids


def test_summary_route_forced_to_company(env):
    """Boshqa korxona route'i so'ralsa ham o'z route'iga kesiladi."""
    st = env
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, data, _ = _request(env, "/api/summary?route=r2", headers=hdr)
    assert status == 200
    assert data.get("ok") is True


def test_admin_sees_all_routes(env):
    st = env
    st.dashboard_user_add("admin", *server_mod._hash_password("to'g'ri"),
                          role="ADMIN")
    _, _, cookies = _login(env, "admin", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, data, _ = _request(env, "/api/routes", headers=hdr)
    ids = {r.get("id") for r in data["routes"]}
    assert {"r1", "r2"} <= ids


def test_dispatcher_blocked_from_admin_endpoints(env):
    """DISPATCHER /api/users (ADMIN-only) dan 403 oladi."""
    st = env
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, _, _ = _request(env, "/api/users", headers=hdr)
    assert status == 403


def test_director_blocked_from_myai(env):
    """DIRECTOR /api/myai (ADMIN-only) dan 403 oladi."""
    st = env
    st.dashboard_user_add("dir", *server_mod._hash_password("to'g'ri"),
                          role="DIRECTOR", company="XTEST")
    _, _, cookies = _login(env, "dir", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, _, _ = _request(env, "/api/myai/status", headers=hdr)
    assert status == 403


def test_dispatcher_blocked_from_brutto(env):
    """DISPATCHER brutto (moliyaviy) ni ko'rmaydi."""
    st = env
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, _, _ = _request(env, "/api/brutto", headers=hdr)
    assert status == 403


# ------------------------------------------------------ user administration

def test_admin_creates_user_via_api(env):
    st = env
    st.dashboard_user_add("root", *server_mod._hash_password("to'g'ri"),
                          role="ADMIN")
    _, _, cookies = _login(env, "root", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}",
           "Content-Type": "application/json"}
    body = json.dumps({"action": "create", "username": "meneger",
                       "role": "MANAGER", "company": "XTEST",
                       "password": "1234"}).encode()
    status, data, _ = _request(env, "/api/dashboard-users", method="POST",
                               headers={**hdr, "Content-Length": str(len(body))},
                               body=body)
    assert status == 200 and data.get("ok") is True
    row = st.dashboard_user_get(username="meneger")
    assert row["role"] == "MANAGER"
    assert row["company"] == "XTEST"


def test_nonadmin_cannot_manage_users(env):
    st = env
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}",
           "Content-Length": "2", "Content-Type": "application/json"}
    status, _, _ = _request(env, "/api/dashboard-users", method="POST",
                            headers=hdr, body=b"{}")
    assert status == 403


def test_admin_list_users_scrubs_secrets(env):
    st = env
    st.dashboard_user_add("root", *server_mod._hash_password("to'g'ri"),
                          role="ADMIN")
    uid = st.dashboard_user_add("ali", *server_mod._hash_password("x"),
                                role="DISPATCHER", company="XTEST")
    st.dashboard_user_update(uid, active=0)
    _, _, cookies = _login(env, "root", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"
                     f"; bm_token=abc"}
    status, data, _ = _request(env, "/api/dashboard-users?action=list",
                               headers=hdr)
    assert status == 200
    users = {u["username"]: u for u in data["users"]}
    assert "ali" in users
    assert users["ali"]["active"] is False
    assert users["ali"]["company"] == "XTEST"
    for u in data["users"]:
        assert "password_hash" not in u
        assert "salt" not in u


def test_admin_edits_user_role_and_company(env):
    """TAHRIRLASH: rol va kompaniya almashadi (parol berilmasa o'zgarmaydi)."""
    st = env
    st.dashboard_user_add("root", *server_mod._hash_password("to'g'ri"),
                          role="ADMIN")
    uid = st.dashboard_user_add("ali", *server_mod._hash_password("eski"),
                                role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "root", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}",
           "Content-Type": "application/json"}
    body = json.dumps({"action": "update", "username": "ali",
                       "role": "MANAGER", "company": "YTEST"}).encode()
    status, data, _ = _request(env, "/api/dashboard-users", method="POST",
                               headers={**hdr, "Content-Length": str(len(body))},
                               body=body)
    assert status == 200 and data.get("ok") is True
    row = st.dashboard_user_get(row_id=uid)
    assert row["role"] == "MANAGER"
    assert row["company"] == "YTEST"
    # Parol kiritilmagani uchun eski parol o'z kuchini saqlaydi.
    _, od, _ = _login(env, "ali", "eski")
    assert od.get("ok") is True


def test_admin_deletes_user_by_username(env):
    """O'CHIRISH tugmasi faqat username yuboradi — server username bilan topadi."""
    st = env
    st.dashboard_user_add("root", *server_mod._hash_password("to'g'ri"),
                          role="ADMIN")
    uid = st.dashboard_user_add("ali", *server_mod._hash_password("x"),
                                role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "root", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}",
           "Content-Type": "application/json"}
    body = json.dumps({"action": "delete", "username": "ali"}).encode()
    status, data, _ = _request(env, "/api/dashboard-users", method="POST",
                               headers={**hdr, "Content-Length": str(len(body))},
                               body=body)
    assert status == 200 and data.get("ok") is True
    assert st.dashboard_user_get(row_id=uid) is None
    # O'zini o'chirish taqiqlangan.
    body2 = json.dumps({"action": "delete", "username": "root"}).encode()
    status, data, _ = _request(env, "/api/dashboard-users", method="POST",
                               headers={**hdr, "Content-Length": str(len(body2))},
                               body=body2)
    assert data.get("ok") is False
    assert st.dashboard_user_get(username="root") is not None


def test_admin_resets_password_only_updates_password(env):
    """PAROL tugmasi: faqat parol yuboriladi — rol/kompaniya o'zgarmaydi."""
    st = env
    st.dashboard_user_add("root", *server_mod._hash_password("to'g'ri"),
                          role="ADMIN")
    st.dashboard_user_add("ali", *server_mod._hash_password("eski"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "root", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}",
           "Content-Type": "application/json"}
    body = json.dumps({"action": "update", "username": "ali",
                       "password": "yangi-1234"}).encode()
    status, data, _ = _request(env, "/api/dashboard-users", method="POST",
                               headers={**hdr, "Content-Length": str(len(body))},
                               body=body)
    assert status == 200 and data.get("ok") is True
    row = st.dashboard_user_get(username="ali")
    assert row["role"] == "DISPATCHER"
    assert row["company"] == "XTEST"
    # Yangi parol kiradi, eski parol kirmaydi.
    _, nd, _ = _login(env, "ali", "yangi-1234")
    assert nd.get("ok") is True
    _, od, _ = _login(env, "ali", "eski")
    assert od.get("ok") is False


def test_admin_toggles_active_only_keeps_company(env):
    """Faol checkbox: faqat active yuboriladi — kompaniya saqlanadi."""
    st = env
    st.dashboard_user_add("root", *server_mod._hash_password("to'g'ri"),
                          role="ADMIN")
    uid = st.dashboard_user_add("ali", *server_mod._hash_password("x"),
                                role="DIRECTOR", company="XTEST")
    _, _, cookies = _login(env, "root", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}",
           "Content-Type": "application/json"}
    body = json.dumps({"action": "update", "username": "ali",
                       "active": False}).encode()
    status, data, _ = _request(env, "/api/dashboard-users", method="POST",
                               headers={**hdr, "Content-Length": str(len(body))},
                               body=body)
    assert status == 200 and data.get("ok") is True
    row = st.dashboard_user_get(row_id=uid)
    assert row["active"] == 0
    assert row["role"] == "DIRECTOR"
    assert row["company"] == "XTEST"


# ------------------------------------------------- rol/xavfsizlik hardending

def test_viewer_can_login_but_only_me(env):
    """VIEWER /api/me ni ko'radi, boshqa endpointlardan 403 oladi."""
    st = env
    st.dashboard_user_add("viewer", *server_mod._hash_password("to'g'ri"),
                          role="VIEWER", company="XTEST")
    _, _, cookies = _login(env, "viewer", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, data, _ = _request(env, "/api/me", headers=hdr)
    assert status == 200 and data["ok"] is True
    assert data["role"] == "VIEWER"
    status, _, _ = _request(env, "/api/summary", headers=hdr)
    assert status == 403


def test_dispatcher_blocked_from_sms_log(env):
    st = env
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, _, _ = _request(env, "/api/sms-log", headers=hdr)
    assert status == 403


def test_dispatcher_blocked_from_sms_routes(env):
    st = env
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, _, _ = _request(env, "/api/sms/routes", headers=hdr)
    assert status == 403


def test_dispatcher_cannot_change_electricity_price(env):
    """Elektrik narxini ko'rish DISPATCHER'ga, o'zgartirish MANAGER'ga."""
    st = env
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}",
           "Content-Type": "application/json"}
    status, _, _ = _request(env, "/api/electricity/price", headers=hdr)
    assert status == 200
    body = json.dumps({"rate": 10}).encode()
    status, _, _ = _request(env, "/api/electricity/price", method="POST",
                            headers={**hdr, "Content-Length": str(len(body))},
                            body=body)
    assert status == 403


def test_dispatcher_blocked_from_telegram_user_admin(env):
    """/api/users/* POST (rol/delete/link) faqat ADMIN uchun."""
    st = env
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}",
           "Content-Type": "application/json"}
    body = json.dumps({"role": "ADMIN"}).encode()
    status, _, _ = _request(env, "/api/users/5/role", method="POST",
                            headers={**hdr, "Content-Length": str(len(body))},
                            body=body)
    assert status == 403


def test_dispatcher_blocked_from_dispatcher_routing(env):
    """Dispetcher route/telefon konfiguratsiyasi faqat ADMIN uchun."""
    st = env
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}",
           "Content-Type": "application/json"}
    body = json.dumps({"routes": []}).encode()
    status, _, _ = _request(env, "/api/dispatchers/5/routes", method="POST",
                            headers={**hdr, "Content-Length": str(len(body))},
                            body=body)
    assert status == 403


def test_director_sms_log_scoped(env, monkeypatch):
    """DIRECTOR faqat o'z korxonasi route'idagi SMS jurnalini ko'radi."""
    from bm_automation.app.notifications import sms_notify as sn
    rows = [
        {"id": 1, "route_id": "r2", "route_name": "20-yo'nalish",
         "phone": "1", "status": "DELIVERED", "name": "A", "message": "m"},
        {"id": 2, "route_id": "r1", "route_name": "10-yo'nalish",
         "phone": "2", "status": "DELIVERED", "name": "B", "message": "m"},
    ]
    monkeypatch.setattr(sn, "sms_log", lambda *a, **k: {
        "rows": rows, "count": 2,
        "counts": {"DELIVERED": 2, "total": 2}, "configured": True})
    st = env
    st.dashboard_user_add("dir", *server_mod._hash_password("to'g'ri"),
                          role="DIRECTOR", company="XTEST")
    _, _, cookies = _login(env, "dir", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, data, _ = _request(env, "/api/sms-log", headers=hdr)
    assert status == 200 and data["ok"] is True
    assert [r["id"] for r in data["rows"]] == [2]
    assert data["count"] == 1
    assert data["counts"]["total"] == 1
    # Admin esa hammasini ko'radi.
    st.dashboard_user_add("root", *server_mod._hash_password("to'g'ri"),
                          role="ADMIN")
    _, _, cookies = _login(env, "root", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, data, _ = _request(env, "/api/sms-log", headers=hdr)
    assert status == 200
    assert len(data["rows"]) == 2
    assert data["count"] == 2


def test_director_sms_routes_scoped(env):
    """DIRECTOR SMS yo'nalishlar ro'yxatida faqat o'z korxonasini ko'radi."""
    st = env
    st.dashboard_user_add("dir", *server_mod._hash_password("to'g'ri"),
                          role="DIRECTOR", company="XTEST")
    _, _, cookies = _login(env, "dir", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, data, _ = _request(env, "/api/sms/routes", headers=hdr)
    assert status == 200 and data["ok"] is True
    ids = [r["route_id"] for r in data["routes"]]
    assert ids == ["r1"]


class _FakeDocMgr:
    def __init__(self, docs):
        self._docs = docs

    def list_documents(self, category="", status="", search="", driver_id=""):
        return [dict(d) for d in self._docs]

    def get_document(self, doc_id):
        return next((dict(d) for d in self._docs if d["id"] == doc_id), None)


def test_director_documents_scoped_by_driver_route(env, monkeypatch):
    """DIRECTOR faqat o'z route'idagi haydovchilar xujjatlarini ko'radi."""
    from bm_automation.app.documents import manager as mgr_mod
    docs = [
        {"id": 1, "title": "Shartnoma-1", "driver_id": "d1"},
        {"id": 2, "title": "Shartnoma-2", "driver_id": "d2"},
    ]
    monkeypatch.setattr(mgr_mod, "get_manager",
                        lambda: _FakeDocMgr(docs))
    st = env
    st.save_driver("d1", full_name="Ali", route_id="r1")
    st.save_driver("d2", full_name="Vali", route_id="r2")
    st.dashboard_user_add("dir", *server_mod._hash_password("to'g'ri"),
                          role="DIRECTOR", company="XTEST")
    _, _, cookies = _login(env, "dir", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, data, _ = _request(env, "/api/documents", headers=hdr)
    assert status == 200 and data["ok"] is True
    assert [d["id"] for d in data["documents"]] == [1]
    # Tafsilot: o'z route'idagi xujjat ochiladi, boshqasi 404 beradi.
    status, _, _ = _request(env, "/api/documents/1", headers=hdr)
    assert status == 200
    status, _, _ = _request(env, "/api/documents/2", headers=hdr)
    assert status == 404
    # Admin ikkalasini ham ko'radi.
    st.dashboard_user_add("root", *server_mod._hash_password("to'g'ri"),
                          role="ADMIN")
    _, _, cookies = _login(env, "root", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, data, _ = _request(env, "/api/documents", headers=hdr)
    assert status == 200
    assert [d["id"] for d in data["documents"]] == [1, 2]


# ------------------------------------------- haydovchi/avtobus PII qamrovi

def test_driver_detail_scoped_to_route(env, monkeypatch):
    """/api/drivers/{id} boshqa route haydovchisi uchun 404 beradi."""
    from bm_automation.app.dashboard.server import DashboardHandler
    called = []

    def fake_detail(self, driver_id, qs):
        called.append(driver_id)
        return {"ok": True, "driver_id": driver_id}

    monkeypatch.setattr(DashboardHandler, "_driver_detail", fake_detail)
    st = env
    st.save_driver("d1", full_name="Ali", route_id="r1")
    st.save_driver("d2", full_name="Vali", route_id="r2")
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, _, _ = _request(env, "/api/drivers/d1", headers=hdr)
    assert status == 200
    assert called == ["d1"]
    status, _, _ = _request(env, "/api/drivers/d2", headers=hdr)
    assert status == 404
    assert called == ["d1"]
    # Admin ikkalasini ham ocha oladi.
    st.dashboard_user_add("root", *server_mod._hash_password("to'g'ri"),
                          role="ADMIN")
    _, _, cookies = _login(env, "root", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, _, _ = _request(env, "/api/drivers/d2", headers=hdr)
    assert status == 200
    assert called == ["d1", "d2"]


def test_drivers_list_scoped(env):
    """/api/drivers/list faqat o'z route'i haydovchilarini qaytaradi."""
    st = env
    st.save_driver("d1", full_name="AAA", route_id="r1")
    st.save_driver("d2", full_name="BBB", route_id="r2")
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, data, _ = _request(env, "/api/drivers/list", headers=hdr)
    assert status == 200 and data["ok"] is True
    assert [d["driver_id"] for d in data["drivers"]] == ["d1"]
    st.dashboard_user_add("root", *server_mod._hash_password("to'g'ri"),
                          role="ADMIN")
    _, _, cookies = _login(env, "root", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, data, _ = _request(env, "/api/drivers/list", headers=hdr)
    assert {d["driver_id"] for d in data["drivers"]} == {"d1", "d2"}


def test_driver_write_scoped_to_route(env):
    """/api/drivers/{id} POST (profil/rasm) boshqa route uchun 404."""
    st = env
    st.save_driver("d1", full_name="Ali", route_id="r1")
    st.save_driver("d2", full_name="Vali", route_id="r2")
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}",
           "Content-Type": "application/json"}
    body = json.dumps({"phone": "+998"}).encode()
    status, _, _ = _request(env, "/api/drivers/d2/profile", method="POST",
                            headers={**hdr, "Content-Length": str(len(body))},
                            body=body)
    assert status == 404


def test_vehicle_detail_scoped(env, monkeypatch):
    """/api/vehicles/{id} boshqa route avtobusi uchun 404 beradi."""
    from bm_automation.app.dashboard import metrics as metrics_mod
    monkeypatch.setattr(metrics_mod.Metrics, "vehicle_detail",
                        lambda self, vid, f: {"ok": True, "id": vid})
    st = env
    st.save_vehicle("v1", plate_number="01A", route_id="r1")
    st.save_vehicle("v2", plate_number="02B", route_id="r2")
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, data, _ = _request(env, "/api/vehicles/v1", headers=hdr)
    assert status == 200 and data["ok"] is True
    status, _, _ = _request(env, "/api/vehicles/v2", headers=hdr)
    assert status == 404


def test_appeals_scoped(env, monkeypatch):
    """/api/appeals ro'yxat va o'chirish o'z route'i bilan cheklanadi."""
    from bm_automation.app.dashboard import metrics as metrics_mod
    monkeypatch.setattr(metrics_mod.Metrics, "appeals_report", lambda self, f: {
        "filters": dict(f or {}),
        "appeals": [
            {"id": 1, "driver_id": "d1", "driver_name": "Ali",
             "title": "T1", "text": "", "status": "YANGI", "reply": "",
             "replied_at": "", "created_at": "", "updated_at": ""},
            {"id": 2, "driver_id": "d2", "driver_name": "Vali",
             "title": "T2", "text": "", "status": "YANGI", "reply": "",
             "replied_at": "", "created_at": "", "updated_at": ""},
        ],
        "total": 2,
    })
    st = env
    st.save_driver("d1", full_name="Ali", route_id="r1")
    st.save_driver("d2", full_name="Vali", route_id="r2")
    aid1 = st.appeal_add("d1", title="T1")
    aid2 = st.appeal_add("d2", title="T2")
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, data, _ = _request(env, "/api/appeals", headers=hdr)
    assert status == 200
    assert [a["id"] for a in data["appeals"]] == [1]
    # Boshqa route murojaatini o'chirish mumkin emas.
    body = json.dumps({"action": "delete", "id": aid2}).encode()
    status, data, _ = _request(env, "/api/appeals", method="POST",
                               headers={**hdr,
                                        "Content-Length": str(len(body)),
                                        "Content-Type": "application/json"},
                               body=body)
    assert status == 404 and data.get("ok") is False
    # O'z route'idagini o'chirish mumkin.
    body = json.dumps({"action": "delete", "id": aid1}).encode()
    status, data, _ = _request(env, "/api/appeals", method="POST",
                               headers={**hdr,
                                        "Content-Length": str(len(body)),
                                        "Content-Type": "application/json"},
                               body=body)
    assert status == 200 and data.get("ok") is True
    ids = [r["id"] for r in st.appeals_list()]
    assert ids == [aid2]


# ----------------------------------------------------- hisobotlar qamrovi

def test_salary_report_scoped_to_company(env, monkeypatch):
    """Ish haqi hisoboti faqat o'z korxonasi route'larini qamraydi."""
    from bm_automation.app.dashboard import export as export_mod
    calls = []

    def fake_export(from_date, to_date, route=""):
        calls.append((from_date, to_date, route))
        return b"{}"

    monkeypatch.setattr(export_mod, "salary_export", fake_export)
    st = env
    st.dashboard_user_add("mgr", *server_mod._hash_password("to'g'ri"),
                          role="MANAGER", company="XTEST")
    _, _, cookies = _login(env, "mgr", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    q = "/api/salary?from=2026-01-01&to=2026-01-31"
    status, _, _ = _request(env, q, headers=hdr)
    assert status == 200
    assert calls and calls[0][2] == "r1"
    # DISPATCHER salary'ga kira olmaydi (rang 2).
    st.dashboard_user_add("disp", *server_mod._hash_password("to'g'ri"),
                          role="DISPATCHER", company="XTEST")
    _, _, cookies = _login(env, "disp", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, _, _ = _request(env, q, headers=hdr)
    assert status == 403
    # ADMIN to'liq eksport qiladi (route bo'sh).
    st.dashboard_user_add("root", *server_mod._hash_password("to'g'ri"),
                          role="ADMIN")
    _, _, cookies = _login(env, "root", "to'g'ri")
    hdr = {"Cookie": f"bm_session={cookies['bm_session']}"}
    status, _, _ = _request(env, q, headers=hdr)
    assert status == 200
    assert calls[-1][2] == ""


def test_salary_export_filters_by_route(monkeypatch):
    """salary_export route parametri bilan faqat o'sha yo'nalishlarni oladi."""
    from io import BytesIO

    from bm_automation.app.dashboard import export as export_mod
    from openpyxl import load_workbook

    ALL = [
        {"driver_id": "d1", "name": "Ali", "route_id": "r1", "trips": 5,
         "working_days": 3, "km": 100, "km_rate": 10, "gross_pay": 1000,
         "tax": 120, "fines": 0, "net_pay": 880, "rating": 5,
         "attendance": 100},
        {"driver_id": "d2", "name": "Vali", "route_id": "r2", "trips": 4,
         "working_days": 2, "km": 90, "km_rate": 10, "gross_pay": 900,
         "tax": 108, "fines": 50, "net_pay": 742, "rating": 4,
         "attendance": 90},
    ]

    class FakeMet:
        def drivers(self, f):
            sel = f.get("route") or ""
            ids = {t for t in str(sel).replace(",", " ").split() if t}
            return [d for d in ALL if not ids or d["route_id"] in ids]

        def _companies(self):
            return {}

        def _route_names(self):
            return {}

    monkeypatch.setattr(export_mod.m, "Metrics", FakeMet)
    buf = BytesIO(export_mod.salary_export("2026-01-01", "2026-01-31",
                                           route="r1"))
    assert set(load_workbook(buf).sheetnames) == {"Xulosa", "Noma'lum — r1"}
    buf = BytesIO(export_mod.salary_export("2026-01-01", "2026-01-31"))
    assert set(load_workbook(buf).sheetnames) == {
        "Xulosa", "Noma'lum — r1", "Noma'lum — r2"}